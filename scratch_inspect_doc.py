import docx
import sys

def inspect(docx_path, out_path):
    doc = docx.Document(docx_path)
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write(f"Doc: {docx_path}\n")
        f.write(f"Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}\n\n")
        f.write("=== PARAGRAPHS ===\n")
        for i, p in enumerate(doc.paragraphs):
            txt = p.text.strip()
            if txt:
                f.write(f"[{i}] ({p.style.name}): {txt}\n")
        f.write("\n=== TABLES ===\n")
        for t_idx, table in enumerate(doc.tables):
            f.write(f"\n--- TABLE {t_idx} ({len(table.rows)} rows, {len(table.columns)} cols) ---\n")
            for r_idx, row in enumerate(table.rows):
                row_txt = [c.text.strip().replace('\n', ' ') for c in row.cells]
                f.write(f"  Row {r_idx}: {' | '.join(row_txt)}\n")

if __name__ == '__main__':
    inspect('docs/MoTa_HeThong_QLNL_NexNova_backup_20261013.docx', 'docs/backup_inspect.txt')
    inspect('docs/MoTa_HeThong_QLNL_NexNova.docx', 'docs/current_inspect.txt')
    print("Done inspecting.")
